from __future__ import annotations
from dataclasses import dataclass, field
from pathlib import Path
from collections import defaultdict
from typing import Optional
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

CAPABILITIES = ["Digital Marketing Strategy", "Content Marketing", "Paid Digital Media", "Social Media", "Influencer Marketing", "UX & Web Design", "SEO", "Marketing Automation", "Data & Analytics"]
REGION_MAP = {"DE":"EUR", "GERMANY":"EUR", "TR":"EUR", "TURKEY":"EUR", "TÜRKIYE":"EUR", "IR":"AMEA", "IRAN":"AMEA", "RU":"AMEA", "RUSSIA":"AMEA", "KR":"AMEA", "SOUTH KOREA":"AMEA"}

@dataclass
class SkillScore:
    capability: str; skill: str; current: Optional[float]; target: Optional[float]; include: int; row: int
@dataclass
class Submission:
    source_file: str; template: str; country: str; business_unit: str; assessment_year: int
    skills: list[SkillScore] = field(default_factory=list); issues: list[str] = field(default_factory=list)
    @property
    def region(self): return REGION_MAP.get(self.country.strip().upper(), "UNMAPPED")
    def capability_scores(self, field_name: str):
        groups=defaultdict(list)
        for s in self.skills:
            if s.include == 1:
                value=getattr(s,field_name)
                if value is None: groups[s.capability].append(None)
                else: groups[s.capability].append(value)
        out={}
        for cap,values in groups.items(): out[cap]=None if any(v is None for v in values) else round(sum(values)/len(values),2)
        return out
    def overall(self, field_name: str):
        values=self.capability_scores(field_name)
        present=[values.get(c) for c in CAPABILITIES]
        return None if any(v is None for v in present) else round(sum(present)/9,2)

def num(value):
    try:
        v=float(value)
        return v if 0 <= v <= 4 and v.is_integer() else None
    except (TypeError,ValueError): return None

def detect_template(ws):
    headers=[str(ws.cell(4,c).value or '') for c in range(1,25)]
    all_text=' '.join(str(cell.value or '') for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,40),values_only=False) for cell in row[:18])
    if 'Reviewer Notes (AMEA)' in headers: return 'amea'
    if 'Google Ads is not available in Iran' in all_text: return 'iran'
    # country-specific templates require a controlled country alias list in production
    return 'europe'

def parse_workbook(path: str|Path, year=2026) -> Submission:
    path=Path(path); wb=load_workbook(path,read_only=True,data_only=True,keep_vba=path.suffix.lower()=='.xlsm')
    required={'Sign-up','DMMM Questionnaire','Summary outcome'}
    missing=required-set(wb.sheetnames)
    if missing: raise ValueError(f'{path.name}: missing sheets: {", ".join(sorted(missing))}')
    signup=wb['Sign-up']; q=wb['DMMM Questionnaire']
    country=str(signup['D7'].value or '').strip(); bu=str(signup['D9'].value or '').strip()
    submission=Submission(path.name,detect_template(q),country,bu,year)
    if not country: submission.issues.append('Missing Market/Country in Sign-up!D7')
    if not bu: submission.issues.append('Missing Business Unit in Sign-up!D9')
    capability=''
    for r in range(5,q.max_row+1):
        raw_cap=q.cell(r,2).value; skill=q.cell(r,3).value
        if raw_cap and not str(raw_cap).startswith('='): capability=str(raw_cap).strip()
        if not skill: continue
        include=q.cell(r,16).value
        include=1 if include in (None,'',1,'1',True) else 0 if include in (0,'0',False) else -1
        if include == -1:
            submission.issues.append(f'Row {r}: invalid inclusion flag P={include!r}'); continue
        current=num(q.cell(r,11).value); target=num(q.cell(r,12).value)
        if include and q.cell(r,11).value not in (None,'') and current is None: submission.issues.append(f'Row {r}: invalid current score')
        if include and q.cell(r,12).value not in (None,'') and target is None: submission.issues.append(f'Row {r}: invalid target score')
        submission.skills.append(SkillScore(capability,str(skill).strip(),current,target,include,r))
    return submission

def parse_batch(folder: str|Path):
    files=sorted(list(Path(folder).glob('*.xlsx'))+list(Path(folder).glob('*.xlsm')))
    records=[]; errors=[]; keys=set()
    for f in files:
        try:
            s=parse_workbook(f); key=(s.country.upper(),s.business_unit.upper(),s.assessment_year)
            if key in keys: s.issues.append('Duplicate country/BU/year submission in this batch')
            keys.add(key); records.append(s)
        except Exception as e: errors.append((f.name,str(e)))
    return records,errors

def build_tracker(submissions:list[Submission], output:str|Path):
    wb=Workbook(); ws=wb.active; ws.title='Results'
    headers=['Region','BU','Market','Assessment Year','Template','Current Self Score','Validated Current Score','Target +1Y','Current Completion','Validation Status']+CAPABILITIES
    ws.append(headers)
    for s in submissions:
        cap=s.capability_scores('current'); cur=s.overall('current'); tgt=s.overall('target')
        completion='Complete' if cur is not None else 'Blocked: missing included current score'
        validation='Pending Global Data' if cur is not None else 'Not ready'
        ws.append([s.region,s.business_unit,s.country,s.assessment_year,s.template,cur,None,tgt,completion,validation]+[cap.get(c) for c in CAPABILITIES])
    _style(ws)
    for region in sorted({s.region for s in submissions}):
        sh=wb.create_sheet(f'{region} REGION')
        sh.append(headers)
        for row in ws.iter_rows(min_row=2,values_only=True):
            if row[0]==region: sh.append(list(row))
        _style(sh)
    log=wb.create_sheet('Intake Log'); log.append(['Source File','Region','BU','Market','Template','Skills Parsed','Issues'])
    for s in submissions: log.append([s.source_file,s.region,s.business_unit,s.country,s.template,len(s.skills),' | '.join(s.issues) or 'OK'])
    _style(log)
    queue=wb.create_sheet('Validation Queue'); queue.append(['Region','BU','Market','Capability','Skill','Self Current','GDM Proposed Current','Source Metric','Threshold Version','Global Verification','Decision'])
    for s in submissions:
        for skill in s.skills:
            if skill.include: queue.append([s.region,s.business_unit,s.country,skill.capability,skill.skill,skill.current,None,None,'2026.1','Pending','Pending'])
    _style(queue)
    Path(output).parent.mkdir(parents=True,exist_ok=True); wb.save(output)

def _style(ws):
    fill=PatternFill('solid',fgColor='006B3F')
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=fill; cell.alignment=Alignment(wrap_text=True)
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        letter=col[0].column_letter; ws.column_dimensions[letter].width=min(max(max(len(str(c.value or '')) for c in col)+2,12),34)
