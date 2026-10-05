import io
import re
import hashlib
from typing import Dict, Any, List, Optional, Tuple
import openpyxl

CANONICAL_CAPABILITIES = {
    "DMS": {"name": "Digital Marketing Strategy", "rows": range(5, 14)},
    "CONTENT": {"name": "Content Marketing", "rows": range(14, 25)},
    "PAID_MEDIA": {"name": "Paid Digital Media", "rows": range(25, 37)},
    "SOCIAL": {"name": "Social Media", "rows": range(37, 47)},
    "INFLUENCER": {"name": "Influencer Marketing & Partnership", "rows": range(47, 52)},
    "UX": {"name": "User Experience (UX) Web Design / Mobile", "rows": range(52, 61)},
    "SEO": {"name": "Search Engine Optimization (SEO)", "rows": range(61, 72)},
    "DATA_ANALYTICS": {"name": "Data & Analytics", "rows": range(72, 81)},
    "MARKETING_AUTO": {"name": "Marketing Automation", "rows": range(81, 91)}
}

PREPOPULATED_ROWS = {13, 23, 24, 46, 60, 69, 70, 71, 89, 90}

class DMMMEngine:
    @staticmethod
    def normalize_score(raw_val: Any) -> Tuple[Optional[float], str]:
        if raw_val is None or str(raw_val).strip() == "":
            return None, "missing"
        val_str = str(raw_val).strip()
        match = re.match(r"^([0-4])(?:\s*-\s*|\s*\()?", val_str)
        if match:
            return float(match.group(1)), "valid"
        return None, "invalid"

    @classmethod
    def parse_and_calculate(cls, file_bytes: bytes, filename: str = "upload.xlsm") -> Dict[str, Any]:
        file_hash = hashlib.sha256(file_bytes).hexdigest()
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, keep_vba=False)
        except Exception as e:
            return {
                "status": "Blocked by file error",
                "errors": [f"Corrupt or unreadable Excel file: {str(e)}"],
                "checksum": file_hash
            }

        validation_errors: List[str] = []
        validation_warnings: List[str] = []

        if "DMMM Questionnaire" not in wb.sheetnames:
            return {
                "status": "Blocked by template mismatch",
                "errors": ["Missing required sheet: 'DMMM Questionnaire'"],
                "checksum": file_hash
            }

        q_ws = wb["DMMM Questionnaire"]
        r4_val = str(q_ws["R4"].value or "")
        profile = "DMMM_2026_AMEA" if "Reviewer Notes (AMEA)" in r4_val else "DMMM_2026_GLOBAL_WORKBOOK_V1"

        metadata = {}
        if "Sign-up" in wb.sheetnames:
            su_ws = wb["Sign-up"]
            metadata = {
                "submitter": su_ws["D5"].value,
                "country": su_ws["D7"].value,
                "business_unit": su_ws["D9"].value,
                "assessment_date": str(su_ws["D11"].value or ""),
                "review_date": str(su_ws["D13"].value or "")
            }
        else:
            validation_warnings.append("'Sign-up' sheet missing. Country metadata unassigned.")

        skills_detail = []
        capability_scores = {}

        for cap_key, cap_info in CANONICAL_CAPABILITIES.items():
            inc_curr_scores = []
            inc_tgt_scores = []
            total_included = 0

            for row in cap_info["rows"]:
                skill_name = q_ws.cell(row=row, column=3).value
                raw_k = q_ws.cell(row=row, column=11).value
                raw_l = q_ws.cell(row=row, column=12).value
                raw_p = q_ws.cell(row=row, column=16).value
                comments = q_ws.cell(row=row, column=13).value

                try:
                    p_val = int(raw_p) if raw_p is not None else 1
                except ValueError:
                    p_val = 1
                    validation_warnings.append(f"Row {row}: Invalid inclusion flag '{raw_p}', defaulted to 1.")

                is_included = (p_val == 1)
                norm_curr, curr_stat = cls.normalize_score(raw_k)
                norm_tgt, tgt_stat = cls.normalize_score(raw_l)
                is_prepop = row in PREPOPULATED_ROWS

                skills_detail.append({
                    "row": row,
                    "capability": cap_info["name"],
                    "skill_name": skill_name,
                    "is_included": is_included,
                    "is_prepopulated": is_prepop,
                    "raw_self": raw_k,
                    "normalized_self": norm_curr,
                    "raw_target": raw_l,
                    "normalized_target": norm_tgt,
                    "status": "valid" if (is_included and norm_curr is not None) else ("excluded" if not is_included else curr_stat),
                    "comments": comments
                })

                if is_included:
                    total_included += 1
                    if norm_curr is not None:
                        inc_curr_scores.append(norm_curr)
                    else:
                        validation_errors.append(f"Row {row} ({skill_name}): Required score missing/invalid.")

                    if norm_tgt is not None:
                        inc_tgt_scores.append(norm_tgt)

            cap_curr = (sum(inc_curr_scores) / len(inc_curr_scores)) if inc_curr_scores else None
            cap_tgt = (sum(inc_tgt_scores) / len(inc_tgt_scores)) if inc_tgt_scores else None

            capability_scores[cap_key] = {
                "capability_name": cap_info["name"],
                "included_skills_count": total_included,
                "valid_scores_count": len(inc_curr_scores),
                "current_score": round(cap_curr, 2) if cap_curr is not None else None,
                "target_score": round(cap_tgt, 2) if cap_tgt is not None else None,
                "gap": round(cap_tgt - cap_curr, 2) if (cap_curr is not None and cap_tgt is not None) else None
            }

        valid_caps = [c["current_score"] for c in capability_scores.values() if c["current_score"] is not None]
        valid_targets = [c["target_score"] for c in capability_scores.values() if c["target_score"] is not None]

        overall_curr = round(sum(valid_caps) / 9.0, 2) if len(valid_caps) == 9 else None
        overall_tgt = round(sum(valid_targets) / 9.0, 2) if len(valid_targets) == 9 else None

        status = "Ready to calculate" if not validation_errors else "Blocked by input error"

        return {
            "metadata": {
                "file_name": filename,
                "sha256": file_hash,
                "detected_profile": profile,
                "country_info": metadata
            },
            "status": status,
            "validation": {
                "errors": validation_errors,
                "warnings": validation_warnings
            },
            "summary": {
                "overall_current_score": overall_curr,
                "overall_target_score": overall_tgt,
                "gap": round(overall_tgt - overall_curr, 2) if (overall_curr and overall_tgt) else None,
                "is_final": (len(validation_errors) == 0 and len(valid_caps) == 9)
            },
            "capability_breakdown": capability_scores,
            "skills_detail": skills_detail
        }
