from pathlib import Path
from openpyxl import Workbook, load_workbook
from dmmm_engine.core import parse_workbook, parse_batch, build_tracker

def make_book(path,country='DE',bu='Seeds',include=1,current=2,target=3):
 w=Workbook(); su=w.active; su.title='Sign-up'; su['D7']=country; su['D9']=bu
 q=w.create_sheet('DMMM Questionnaire'); q.append([]);q.append([]);q.append([]); q.append(['','Capability','Skill','','','','','','','','Self','Target','','','','Weighting'])
 q.append(['','Digital Marketing Strategy','Vision','','','','','','','',current,target,'','','',include])
 for cap in ['Content Marketing','Paid Digital Media','Social Media','Influencer Marketing','UX & Web Design','SEO','Marketing Automation','Data & Analytics']: q.append(['',cap,'Skill','','','','','','','',current,target,'','','',include])
 w.create_sheet('Summary outcome'); w.save(path)

def test_parse_score_and_exclusion(tmp_path):
 p=tmp_path/'de.xlsm'; make_book(p,include=1); s=parse_workbook(p); assert s.overall('current')==2
 p2=tmp_path/'ir.xlsm'; make_book(p2,country='IR',include=0,current=None); s2=parse_workbook(p2); assert s2.overall('current') is None # all capabilities excluded, blocks finalisation

def test_batch_duplicate_and_tracker(tmp_path):
 make_book(tmp_path/'a.xlsx'); make_book(tmp_path/'b.xlsx'); records,errors=parse_batch(tmp_path); assert not errors and 'Duplicate' in records[1].issues[0]
 out=tmp_path/'tracker.xlsx'; build_tracker(records,out); assert out.exists(); assert 'Results' in load_workbook(out).sheetnames
